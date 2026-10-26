# -*- coding: utf-8 -*-
"""
一触青绿 设计方案 V2 —— 带图版
模板骨架：概述(1.1/1.2) / 2 项目设计思路 / 3 资源需求(软硬件选型) /
          4 建设方案(4.1~4.6) / 5 功能实现逻辑 / 6 术语和缩略语
正文：宋体 / 小四(12pt) / 1.5 倍行距 / 首行缩进 2 字
"""
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
import matplotlib.font_manager as fm

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ----------------------------------------------------------------------------
# 0. 字体与绘图基础
# ----------------------------------------------------------------------------
FONT_PATH = r"C:/Windows/Fonts/simhei.ttf"
fm.fontManager.addfont(FONT_PATH)
_FP = fm.FontProperties(fname=FONT_PATH)
plt.rcParams["font.family"] = _FP.get_name()
plt.rcParams["axes.unicode_minus"] = False

DIAG_DIR = os.path.join(os.path.dirname(__file__), "diagrams")
os.makedirs(DIAG_DIR, exist_ok=True)

# 青绿配色
C_BG = "#F2F8F5"
C_BAND = "#E3F1EC"
C_MAIN = "#1F8A70"
C_MAIN_D = "#15654F"
C_ACC = "#17A2B8"
C_TXT = "#12302A"
C_LINE = "#3A5A52"
C_WHITE = "#FFFFFF"


def _box(ax, x, y, w, h, text, fc="#FFFFFF", ec=C_MAIN, tc=C_TXT,
         fs=10, bold=True, lw=1.3, rad=0.08, z=3):
    ax.add_patch(FancyBboxPatch(
        (x, y), w, h, boxstyle=f"round,pad=0,rounding_size={rad}",
        fc=fc, ec=ec, lw=lw, zorder=z - 1))
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center",
            fontsize=fs, color=tc, fontproperties=_FP,
            fontweight=("bold" if bold else "normal"), zorder=z)


def _arrow(ax, x1, y1, x2, y2, color=C_LINE, lw=1.6, rad=0.0, style="-|>"):
    ax.add_patch(FancyArrowPatch(
        (x1, y1), (x2, y2), arrowstyle=style, mutation_scale=12,
        color=color, lw=lw,
        connectionstyle=f"arc3,rad={rad}", zorder=1))


def _title(ax, text, W):
    ax.text(W / 2, ax.get_ylim()[1] - 0.45, text, ha="center", va="center",
            fontsize=13, color=C_MAIN_D, fontproperties=_FP, fontweight="bold")


def _newfig(W, H):
    fig = plt.figure(figsize=(W, H), dpi=200)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, W)
    ax.set_ylim(0, H)
    ax.axis("off")
    return fig, ax


def _save(fig, name):
    path = os.path.join(DIAG_DIR, name)
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


# ----------------------------------------------------------------------------
# 1. 设计理念图（第 2 章）
# ----------------------------------------------------------------------------
def diag_concept():
    W, H = 10, 10
    fig, ax = _newfig(W, H)
    ax.text(W / 2, H - 0.5, "图2  一触青绿 · 四大设计主张", ha="center",
            fontsize=13, color=C_MAIN_D, fontproperties=_FP, fontweight="bold")
    _box(ax, 3.9, 4.3, 2.2, 1.4, "一触青绿", fc=C_MAIN, ec=C_MAIN_D,
         tc=C_WHITE, fs=14)
    nodes = [
        (3.9, 8.0, 2.2, 1.2, "一触即达\n零步骤操作"),
        (3.9, 0.8, 2.2, 1.2, "自洽明暗\n深色换血"),
        (0.6, 4.3, 2.2, 1.4, "物理真实\n液态玻璃·苹果动效"),
        (7.2, 4.3, 2.2, 1.4, "活的设计\n数据驱动视觉"),
    ]
    for (x, y, w, h, t) in nodes:
        _box(ax, x, y, w, h, t)
    _arrow(ax, 5.0, 8.0, 5.0, 5.7)
    _arrow(ax, 5.0, 4.3, 5.0, 2.0)
    _arrow(ax, 3.9, 5.0, 2.8, 5.0, rad=0.0)
    _arrow(ax, 6.1, 5.0, 7.2, 5.0)
    return _save(fig, "concept.png")


# ----------------------------------------------------------------------------
# 2. 系统架构图（4.1）
# ----------------------------------------------------------------------------
def diag_sys_arch():
    W, H = 10, 14.6
    fig, ax = _newfig(W, H)
    _title(ax, "图4.1  系统总体架构（五层 + 近场/AI）", W)

    lanes = [
        ("用户\n设备", ["HarmonyOS 手机", "平板", "智慧屏", "穿戴设备"]),
        ("端侧应用\n(ArkTS/ArkUI)", ["表现/UI 组件", "动效层\nGlassLight·AppleMotion",
                                  "状态管理", "业务/页面", "数据持久"]),
        ("通信层", ["HTTPS / REST\n(Token 鉴权)", "软总线 SoftBus\n(近场同步)"]),
        ("云端服务\n(RuoYi-Vue)", ["API 网关/鉴权", "碳业务模块 carbon", "管理端 admin(Vue)"]),
        ("数据存储", ["MySQL 8\n(业务库)", "Redis\n(缓存/会话)"]),
    ]
    y = H - 1.4
    h = 2.2
    gap = 0.35
    prev_bottom = None
    for (cat, items) in lanes:
        # 背景泳道
        ax.add_patch(FancyBboxPatch((0.4, y), 9.2, h,
                     boxstyle="round,pad=0,rounding_size=0.1",
                     fc=C_BG, ec="#C9E3D8", lw=1, zorder=0))
        # 左侧色轨
        ax.add_patch(FancyBboxPatch((0.4, y), 1.7, h,
                     boxstyle="round,pad=0,rounding_size=0.1",
                     fc=C_MAIN, ec=C_MAIN_D, lw=1, zorder=1))
        ax.text(1.25, y + h / 2, cat, ha="center", va="center", fontsize=9.5,
                color=C_WHITE, fontproperties=_FP, fontweight="bold", zorder=2)
        # 内部条目
        n = len(items)
        x0, x1 = 2.4, 9.5
        total = x1 - x0
        if n == 1:
            bw = total
            g = 0
        else:
            g = 0.22
            bw = (total - g * (n - 1)) / n
        for i, it in enumerate(items):
            x = x0 + i * (bw + g)
            _box(ax, x, y + 0.28, bw, h - 0.56, it, fs=8.6, lw=1.1, rad=0.07)
        if prev_bottom is not None:
            _arrow(ax, 5.0, y + h, 5.0, prev_bottom, color=C_LINE, lw=1.4)
        prev_bottom = y
        y -= (h + gap)

    # 右侧旁注
    _box(ax, 9.9, 6.0, 0.0, 0.0, "", fc="none", ec="none")  # noop placeholder
    ax.text(9.85, 8.6, "超级终端协同\n(SoftBus 近场)", ha="left", va="center",
            fontsize=8.5, color=C_ACC, fontproperties=_FP, fontweight="bold")
    ax.annotate("", xy=(9.5, 8.0), xytext=(9.8, 8.6),
                arrowprops=dict(arrowstyle="-|>", color=C_ACC, lw=1.3))
    ax.text(9.85, 3.4, "AI 视觉识别\n(云端推理)", ha="left", va="center",
            fontsize=8.5, color=C_ACC, fontproperties=_FP, fontweight="bold")
    ax.annotate("", xy=(9.5, 4.3), xytext=(9.8, 3.4),
                arrowprops=dict(arrowstyle="-|>", color=C_ACC, lw=1.3))
    return _save(fig, "sys_arch.png")


# ----------------------------------------------------------------------------
# 3. 整体功能架构图（4.2）
# ----------------------------------------------------------------------------
def diag_func_arch():
    W, H = 10, 12
    fig, ax = _newfig(W, H)
    _title(ax, "图4.2  整体功能架构（三大簇 · 十模块）", W)
    # 根
    _box(ax, 3.4, 10.4, 3.2, 1.2, "一触青绿 应用", fc=C_MAIN, ec=C_MAIN_D,
         tc=C_WHITE, fs=13)
    clusters = [
        (0.5, 7.4, 2.9, "碳普惠核心", ["碳记录", "积分任务", "碳惠券包", "看板图表"]),
        (3.55, 7.4, 2.9, "知识与互动", ["碳科普", "AI 识别", "智慧服务"]),
        (6.6, 7.4, 2.9, "账户与协同", ["家庭/多账户", "超级终端同步", "设置"]),
    ]
    for (x, y, w, title, mods) in clusters:
        ax.add_patch(FancyBboxPatch((x, y - 0.2), w, 2.6,
                     boxstyle="round,pad=0,rounding_size=0.1",
                     fc=C_BAND, ec=C_MAIN, lw=1.2, zorder=0))
        ax.text(x + w / 2, y + 2.05, title, ha="center", va="center",
                fontsize=11, color=C_MAIN_D, fontproperties=_FP,
                fontweight="bold", zorder=2)
        m = len(mods)
        mx0, mx1 = x + 0.2, x + w - 0.2
        mbw = (mx1 - mx0 - 0.18 * (m - 1)) / m
        for i, md in enumerate(mods):
            bx = mx0 + i * (mbw + 0.18)
            _box(ax, bx, y, mbw, 1.5, md, fs=9, lw=1.1, rad=0.07)
        _arrow(ax, 5.0, 10.4, x + w / 2, y + 2.4, color=C_LINE, lw=1.3)
    return _save(fig, "func_arch.png")


# ----------------------------------------------------------------------------
# 4. 软件架构图（4.3 端侧六层）
# ----------------------------------------------------------------------------
def diag_sw_arch():
    W, H = 10, 14
    fig, ax = _newfig(W, H)
    _title(ax, "图4.3  端侧软件架构（六层分层）", W)
    layers = [
        ("UI 组件层", ["GlassSurface", "GlassDialog", "Components", "CfCard"]),
        ("动效层", ["AppleMotion（时长/曲线/弹簧）", "GlassLight（液态玻璃光照）"]),
        ("状态管理层", ["AppStorage / @StorageLink", "AccountStore", "@State / @Link / @Watch"]),
        ("业务逻辑层", ["pages（导航页）", "services（用例）", "业务规则"]),
        ("数据持久层", ["Preferences（命名空间隔离）"]),
        ("设备/系统层", ["SoftBus（近场）", "系统能力", "AppClock 时钟校准"]),
    ]
    y = H - 1.4
    h = 1.9
    gap = 0.32
    for (title, items) in layers:
        ax.add_patch(FancyBboxPatch((0.4, y), 9.2, h,
                     boxstyle="round,pad=0,rounding_size=0.1",
                     fc=C_BG, ec="#C9E3D8", lw=1, zorder=0))
        ax.add_patch(FancyBboxPatch((0.4, y), 2.0, h,
                     boxstyle="round,pad=0,rounding_size=0.1",
                     fc=C_MAIN, ec=C_MAIN_D, lw=1, zorder=1))
        ax.text(1.4, y + h / 2, title, ha="center", va="center", fontsize=10,
                color=C_WHITE, fontproperties=_FP, fontweight="bold", zorder=2)
        n = len(items)
        x0, x1 = 2.7, 9.5
        total = x1 - x0
        g = 0.2
        bw = (total - g * (n - 1)) / n
        for i, it in enumerate(items):
            x = x0 + i * (bw + g)
            _box(ax, x, y + 0.3, bw, h - 0.6, it, fs=8.6, lw=1.1, rad=0.07)
        if y + h + gap < H:
            _arrow(ax, 5.0, y, 5.0, y - gap, color=C_LINE, lw=1.4)
        y -= (h + gap)
    return _save(fig, "sw_arch.png")


# ----------------------------------------------------------------------------
# 5. 业务流程图（4.4）
# ----------------------------------------------------------------------------
def diag_biz_flow():
    W, H = 15, 8.5
    fig, ax = _newfig(W, H)
    _title(ax, "图4.4  核心业务流程（注册闸门 → 记录 → 核算 → 兑换 → 核销 → 同步）", W)
    top = [
        (0.6, "注册/登录"),
        (2.6, "注册闸门\n锁死"),
        (4.6, "完善资料"),
        (6.6, "首页看板"),
        (8.6, "碳记录\n录入"),
        (10.6, "减排核算"),
        (12.6, "积分累计"),
    ]
    yT = 5.6
    h = 1.3
    bw = 1.7
    for i, (x, t) in enumerate(top):
        _box(ax, x, yT, bw, h, t, fs=9)
        if i > 0:
            _arrow(ax, x - 0.1, yT + h / 2, x - 0.3 + bw, yT + h / 2,
                   color=C_LINE, lw=1.4)
    # 第二行
    yB = 2.2
    bot = [
        (4.6, "云端同步"),
        (6.6, "线下核销"),
        (8.6, "生成核销码\nPT-"),
        (10.6, "券包兑换"),
        (12.6, "积分商城"),
    ]
    for (x, t) in bot:
        _box(ax, x, yB, bw, h, t, fs=9, fc="#EAF6F2")
    # 连线
    _arrow(ax, 13.45, yT, 13.45, yB + h, color=C_LINE, lw=1.4)      # 积分累计→商城
    _arrow(ax, 12.6 + bw / 2, yT, 11.45, yB + h, rad=-0.15)          # →券包
    _arrow(ax, 11.45, yB + h / 2, 10.6 + 0.1, yB + h / 2, lw=1.4)   # 券包→核销码
    _arrow(ax, 9.45, yB + h / 2, 8.6 + 0.1, yB + h / 2, lw=1.4)     # 核销码→核销
    _arrow(ax, 7.45, yB + h / 2, 6.6 + 0.1, yB + h / 2, lw=1.4)     # 核销→同步
    # 说明
    ax.text(0.6, 0.6, "说明：注册成功后主框架第一帧即锁死交互，待「完善资料」页真实上屏(onShown)才解锁；"
            "积分可走商城或券包两条兑现路径，券包生成 PT- 核销码供线下核销，操作最终同步云端。",
            ha="left", va="center", fontsize=7.8, color=C_TXT,
            fontproperties=_FP, wrap=True)
    return _save(fig, "biz_flow.png")


# ----------------------------------------------------------------------------
# 6. 接口对接图（4.5）
# ----------------------------------------------------------------------------
def diag_interface():
    W, H = 10, 10
    fig, ax = _newfig(W, H)
    _title(ax, "图4.5  接口对接（REST / SoftBus / AI 推理 / 数据存取）", W)
    _box(ax, 3.8, 4.2, 2.4, 1.5, "鸿蒙端 App", fc=C_MAIN, ec=C_MAIN_D,
         tc=C_WHITE, fs=12)
    targets = [
        (0.4, 6.6, 2.6, 1.6, "服务端\nREST/HTTPS", "Token 鉴权\ncarbon API"),
        (0.4, 1.2, 2.6, 1.6, "超级终端\nSoftBus", "近场发现\n数据同步"),
        (7.0, 6.6, 2.6, 1.6, "AI 视觉识别\n(云端推理)", "图像→行为\n结构化"),
        (7.0, 1.2, 2.6, 1.6, "MySQL / Redis", "业务库\n缓存/会话"),
    ]
    for (x, y, w, h, t, sub) in targets:
        _box(ax, x, y, w, h, t, fs=9.5)
        ax.text(x + w / 2, y + 0.35, sub, ha="center", va="center",
                fontsize=7.6, color=C_TXT, fontproperties=_FP)
    _arrow(ax, 3.8, 5.4, 3.0, 6.9, rad=0.1)
    _arrow(ax, 3.8, 4.9, 3.0, 2.6, rad=-0.1)
    _arrow(ax, 6.2, 5.4, 7.0, 6.9, rad=-0.1)
    _arrow(ax, 6.2, 4.9, 7.0, 2.6, rad=0.1)
    ax.text(5.0, 7.7, "HTTPS", ha="center", fontsize=8, color=C_ACC,
            fontproperties=_FP, fontweight="bold")
    ax.text(5.0, 1.4, "SoftBus", ha="center", fontsize=8, color=C_ACC,
            fontproperties=_FP, fontweight="bold")
    ax.text(6.6, 7.7, "AI 推理", ha="center", fontsize=8, color=C_ACC,
            fontproperties=_FP, fontweight="bold")
    ax.text(6.6, 1.4, "JDBC", ha="center", fontsize=8, color=C_ACC,
            fontproperties=_FP, fontweight="bold")
    return _save(fig, "interface.png")


# ----------------------------------------------------------------------------
# 7. 硬件架构图（4.6）
# ----------------------------------------------------------------------------
def diag_hw_arch():
    W, H = 10, 10.5
    fig, ax = _newfig(W, H)
    _title(ax, "图4.6  硬件架构（客户端 · 网络 · 服务端）", W)
    # 客户端
    ax.add_patch(FancyBboxPatch((0.4, 1.0), 2.9, 8.0,
                 boxstyle="round,pad=0,rounding_size=0.12",
                 fc=C_BAND, ec=C_MAIN, lw=1.2, zorder=0))
    ax.text(1.85, 8.4, "客户端设备", ha="center", va="center", fontsize=11,
            color=C_MAIN_D, fontproperties=_FP, fontweight="bold")
    _box(ax, 0.7, 6.6, 2.3, 1.2, "HarmonyOS 手机\nSoC/内存/摄像头", fs=8.4)
    _box(ax, 0.7, 4.9, 2.3, 1.2, "平板 / 智慧屏", fs=8.6)
    _box(ax, 0.7, 3.2, 2.3, 1.2, "穿戴设备", fs=9)
    # 网络
    ax.add_patch(FancyBboxPatch((3.7, 1.0), 2.6, 8.0,
                 boxstyle="round,pad=0,rounding_size=0.12",
                 fc="#EAF6F2", ec=C_ACC, lw=1.2, zorder=0))
    ax.text(5.0, 8.4, "网络通道", ha="center", va="center", fontsize=11,
            color=C_ACC, fontproperties=_FP, fontweight="bold")
    for i, t in enumerate(["Wi-Fi / 5G", "蓝牙 / NFC", "SoftBus 近场"]):
        _box(ax, 3.9, 6.8 - i * 1.7, 2.2, 1.2, t, fs=8.6, ec=C_ACC)
    # 服务端
    ax.add_patch(FancyBboxPatch((6.7, 1.0), 2.9, 8.0,
                 boxstyle="round,pad=0,rounding_size=0.12",
                 fc=C_BAND, ec=C_MAIN, lw=1.2, zorder=0))
    ax.text(8.15, 8.4, "服务端硬件", ha="center", va="center", fontsize=11,
            color=C_MAIN_D, fontproperties=_FP, fontweight="bold")
    _box(ax, 7.0, 6.6, 2.3, 1.2, "应用服务器\n多核CPU/4C8G+", fs=8.4)
    _box(ax, 7.0, 4.9, 2.3, 1.2, "数据库服务器\nMySQL 8", fs=8.6)
    _box(ax, 7.0, 3.2, 2.3, 1.2, "缓存/Redis", fs=9)
    # 连接
    _arrow(ax, 3.3, 5.0, 3.7, 5.0, color=C_LINE, lw=1.6)
    _arrow(ax, 6.3, 5.0, 6.7, 5.0, color=C_LINE, lw=1.6)
    return _save(fig, "hw_arch.png")


DIAGS = {
    "concept": ("图2  一触青绿 · 四大设计主张", diag_concept),
    "sys_arch": ("图4.1  系统总体架构", diag_sys_arch),
    "func_arch": ("图4.2  整体功能架构", diag_func_arch),
    "sw_arch": ("图4.3  端侧软件架构", diag_sw_arch),
    "biz_flow": ("图4.4  核心业务流程", diag_biz_flow),
    "interface": ("图4.5  接口对接关系", diag_interface),
    "hw_arch": ("图4.6  硬件架构", diag_hw_arch),
}


# ============================================================================
# 文档生成
# ============================================================================
def set_eastasia(style_or_run, name):
    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def build_docx():
    doc = Document()
    # 默认正文：宋体 / 小四 / 1.5 / 首行缩进2字
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    set_eastasia(normal, "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.space_after = Pt(6)

    # 标题样式（黑体）
    for sname, sz in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        st = doc.styles[sname]
        st.font.name = "黑体"
        st.font.size = Pt(sz)
        st.font.bold = True
        set_eastasia(st, "黑体")
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.first_line_indent = Pt(0)
        st.paragraph_format.line_spacing = 1.3

    # 页边距
    for sec in doc.sections:
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)

    def H1(t):
        p = doc.add_heading(t, level=1)
        return p

    def H2(t):
        return doc.add_heading(t, level=2)

    def H3(t):
        return doc.add_heading(t, level=3)

    def para(t, align=None, size=12, bold=False, italic=False, color=None):
        p = doc.add_paragraph()
        if align:
            p.alignment = align
        r = p.add_run(t)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        set_eastasia(r, "宋体")
        if color:
            r.font.color.rgb = RGBColor(*color)
        return p

    def figure(key, caption, w=15):
        path = DIAGS[key][1]()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Cm(w))
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        r.font.size = Pt(10.5)
        r.font.italic = True
        set_eastasia(r, "宋体")

    # ---------------- 封面 ----------------
    para("一触青绿", align=WD_ALIGN_PARAGRAPH.CENTER, size=30, bold=True,
         color=(0x1F, 0x8A, 0x70))
    para("个人碳足迹记录与碳积分激励应用 · 设计方案",
         align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True)
    para("（图文版 V2.0）", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    para("")
    para("技术底座：OpenHarmony / ArkTS / ArkUI（API 24）  +  若依 RuoYi-Vue 3.8.7（JDK 11 / MySQL 8 / Redis）",
         align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
    para("编制日期：2026 年 10 月 21 日", align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
    para("")

    # ---------------- 目录（手动） ----------------
    H2("目  录")
    toc = [
        "概述", "　1.1 项目背景", "　1.2 系统建设的必要性",
        "2 项目设计思路", "3 资源需求（软硬件选型）", "4 建设方案",
        "　4.1 系统架构", "　4.2 整体功能架构", "　4.3 软件架构",
        "　4.4 业务流程", "　4.5 接口对接", "　4.6 硬件架构",
        "5 功能实现逻辑", "6 术语和缩略语",
    ]
    for t in toc:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(t)
        r.font.size = Pt(12)
        set_eastasia(r, "宋体")

    doc.add_page_break()

    # ---------------- 概述 ----------------
    H1("概述")
    H2("1.1 项目背景")
    para("在“碳达峰、碳中和”国家战略的宏观背景下，个人层面的低碳行为长期面临“看得见、记不住、无激励”"
         "的困境。日常通勤、节能、 recycling 等减碳行为散落于生活各处，既缺少统一可信的记录入口，"
         "也缺乏将其转化为持续动力的激励机制。")
    para("“一触青绿”由此而生——其理念是“一触之间，青绿相随”：以 OpenHarmony 为底座，"
         "打造一款面向个人用户的碳足迹记录与碳积分激励应用。它把复杂的碳核算封装为“一触即达”的轻量操作，"
         "让用户在几乎无学习成本的前提下，完成低碳行为的记录、积分的累积与权益的兑换，"
         "从而形成“记录—激励—再行动”的正向闭环。同时，作为 HarmonyOS 生态应用，"
         "它天然具备多端协同与近场互通的能力，可延伸至家庭账户与超级终端场景。")

    H2("1.2 系统建设的必要性")
    para("第一，政策与民意需要落地抓手。双碳目标最终要落到每个人的行为改变，而现状是个人侧参与通道稀缺、"
         "感知弱。本系统以“轻量记录 + 即时激励”降低参与门槛，是把宏观战略转化为微观行动的有效载体。")
    para("第二，可信与流畅缺一不可。碳数据一旦涉及积分与权益，就必须做到记录可信、核算透明、核销防抵赖；"
         "而面向大众的应用，又必须在动画与流畅度上达到消费级预期。本系统在端侧引入自研液态玻璃与苹果级动效语言，"
         "在业务侧以结构化因子库与核销码保证严谨，二者缺一不可。")
    para("第三，多端与近场是差异化价值。依托 HarmonyOS 的软总线（SoftBus），应用可在手机、平板、"
         "智慧屏、穿戴之间近场协同，并支持家庭多账户，这是传统单端 App 难以低成本实现的。")
    para("第四，示范与沉淀意义。项目沉淀了“注册闸门零帧锁死”“覆盖层弹窗范式”“深色换血”等"
         "一套可复用的 HarmonyOS 工程方法论，对其他同生态应用具有借鉴价值。")

    # ---------------- 2 项目设计思路 ----------------
    H1("2 项目设计思路")
    para("整体设计围绕“一触青绿”四个字展开，凝练为四大设计主张，并用一套统一的导航与视觉语言贯穿始终。")
    figure("concept", "图2  一触青绿四大设计主张：一触即达、物理真实、活的设计、自洽明暗")
    H3("2.1 一触即达")
    para("以“胶囊 + 圆钮”的导航范式为骨架，把最高频的碳记录动作收敛到一触之内；二级页面通过全局路由"
         "（Navigation）瞬时拉起，避免层层点击。注册成功后的资料完善页以“零帧锁死”机制强制前置，"
         "保证流程不被手快打断。")
    H3("2.2 物理真实")
    para("放弃生硬的补间，改用自研 GlassLight 液态玻璃光照模型（135° 单一主光源、菲涅尔亮边、"
         "α 预合成防脏）与 AppleMotion 动效语言（弹簧回弹、统一时长曲线、按压反馈），"
         "让界面具备接近真实物体的重量感与光泽。")
    H3("2.3 活的设计")
    para("视觉并非静态贴图，而是由数据驱动：积分、徽章、碳情绪会实时改变主色与反馈强度，"
         "让“青绿”随用户的低碳行为一起生长。")
    H3("2.4 自洽明暗")
    para("深色模式不是简单反色，而是通过“深色换血”机制（ThemePalette.applyDark + ThemeMode 推送主题刻度），"
         "让所有静态色值在订阅后整体重建，保证浅色/深色两套材质参数自洽、对比度达标。")

    # ---------------- 3 资源需求 ----------------
    H1("3 资源需求（软硬件选型）")
    para("本节给出开发期与运行期的资源清单，并说明关键选型依据。端侧基于 OpenHarmony API 24，"
         "云端基于若依 RuoYi-Vue 3.8.7 单体架构，整体四端（app / server / admin / db）协同。")
    H3("3.1 开发环境")
    para("鸿蒙端：DevEco Studio（API 24 SDK）、HarmonyOS 模拟器或真机；后端：JDK 11、"
         "Maven 3.x、MySQL 8、Redis 6+；管理前端：Node.js + Vue 生态；数据库脚本以 SQL 提供。"
         "构建使用 hvigor 的 assembleHap 产出 HAP。")
    H3("3.2 运行环境（客户端）")
    para("HarmonyOS 设备需支持 API 24 及以上；AI 视觉识别依赖摄像头与端云协同；"
         "近场同步依赖 SoftBus 能力（同账号多设备）。建议运行内存 4GB 以上以保证液态玻璃多档画质流畅。")
    H3("3.3 运行环境（服务端）")
    para("应用服务器推荐 4 核 8GB 及以上（开发期 2 核 4GB 可运行）；数据库服务器独立部署 MySQL 8；"
         "Redis 用于会话与缓存。网络需开放 8089 端口供端侧调用。")
    H3("3.4 选型依据（摘要）")
    para("OpenHarmony 提供原生分布式与软总线能力，契合多端协同诉求；若依框架成熟、权限与代码生成开箱即用，"
         "可把精力聚焦 carbon 业务；MySQL + Redis 组合兼顾事务与性能；自研 GlassLight / AppleMotion"
         " 则是为弥补通用组件在“高级质感”上的空白而引入。")

    # ---------------- 4 建设方案 ----------------
    H1("4 建设方案")
    H2("4.1 系统架构")
    para("系统自顶向下分为五层：用户与多端设备层、端侧应用层、通信层、云端服务层、数据存储层；"
         "并在通信层之外，通过软总线提供近场协同、通过云端推理提供 AI 视觉识别能力。")
    figure("sys_arch", "图4.1  系统总体架构：设备 / 端侧 / 通信 / 云端 / 数据 五层，叠加近场协同与 AI 识别")

    H2("4.2 整体功能架构")
    para("功能按“碳普惠核心 / 知识与互动 / 账户与协同”三大簇组织，共十个一级模块，"
         "彼此通过统一的状态与数据层解耦。")
    figure("func_arch", "图4.2  整体功能架构：三大簇（碳普惠核心、知识与互动、账户与协同）承载十个模块")

    H2("4.3 软件架构")
    para("端侧采用六层分层：UI 组件层负责表现，动效层统一质感与节奏，状态管理层做响应式与多账户隔离，"
         "业务逻辑层承载页面与用例，数据持久层以 Preferences 命名空间隔离，设备/系统层对接 SoftBus 与系统能力。")
    figure("sw_arch", "图4.3  端侧软件架构：UI / 动效 / 状态 / 业务 / 数据 / 设备 六层及其依赖方向")

    H2("4.4 业务流程")
    para("核心闭环为：注册登录 → 注册闸门锁死 → 完善资料 → 首页看板 → 碳记录录入 → 减排核算 → 积分累计；"
         "积分可分别走向“积分商城”或“碳惠券包”两条兑现路径，券包生成 PT- 核销码供线下核销，"
         "所有操作最终同步云端。")
    figure("biz_flow", "图4.4  核心业务流程：闸门锁死 → 记录核算 → 积分 → 商城/券包 → 核销 → 云端同步")

    H2("4.5 接口对接")
    para("端侧与服务端以 HTTPS / REST 通信并采用 Token 鉴权（carbon 业务 API）；端侧与超级终端以 SoftBus"
         " 近场发现与同步；AI 视觉识别以图像上云推理返回结构化行为；服务端经 JDBC 读写 MySQL、以 Redis 缓存会话。")
    figure("interface", "图4.5  接口对接：REST(Token) / SoftBus(近场) / AI 推理 / JDBC 四类通道")

    H2("4.6 硬件架构")
    para("客户端涵盖手机（含 SoC、内存、摄像头）、平板/智慧屏、穿戴设备；网络通道包含 Wi-Fi/5G、"
         "蓝牙/NFC 与 SoftBus 近场；服务端由应用服务器、数据库服务器与 Redis 缓存构成。")
    figure("hw_arch", "图4.6  硬件架构：客户端设备 — 网络通道 — 服务端硬件 三段式")

    # ---------------- 5 功能实现逻辑 ----------------
    H1("5 功能实现逻辑")
    H3("5.1 注册闸门零帧锁死")
    para("将 profileGate 由响应式状态改为“私有布尔 + 触发器”的零帧锁，读取到 profileSetupIntent 即立刻锁死主框架；"
         "资料页以 NavDestination 的 onShown（系统级“已显示”）作为解锁信号，并辅以 420ms 兜底，"
         "确保“页面真的上屏”后才放开交互，杜绝手快点穿。")
    H3("5.2 碳记录与减排核算")
    para("用户选择低碳行为并录入数量，系统按因子库（行为类型 × 排放因子）计算减排量；"
         "减排量经规则换算为积分，写入账户并驱动看板与徽章刷新。")
    H3("5.3 积分、徽章与碳惠券包")
    para("积分累计后，用户可在积分商城直接兑换，或在碳惠券包中以积分换取权益券；券包兑换成功生成 PT- 前缀的"
         "核销码（含二维码），供线下核销商户扫码核销，核销状态回写并防重复抵赖。")
    H3("5.4 液态玻璃与苹果动效")
    para("GlassLight 以 135° 单一光源建立菲涅尔亮边与 α 预合成，并提供三档画质降级与滚动视差；"
         "AppleMotion 统一全局时长阶梯与曲线，并以 springMotion 弹簧实现回弹与按压反馈，"
         "替代散落的 620ms 手写补间。弹窗采用静态 .animation 避免取消后再开丢失入场（修复三元竞态）。")
    H3("5.5 深色换血与系统时钟")
    para("深色模式通过 ThemeMode 推送主题刻度，令所有静态色值订阅重建；系统时钟统一以 AppClock 基于 Date.now()"
         " 提供 now/today/dateStr，避免设备时钟漂移带来的展示错乱。")
    H3("5.6 多账户与近场同步")
    para("Preferences 以账号命名空间隔离，切换账户即重建数据树；同账号多设备经 SoftBus 近场发现并同步碳记录与积分，"
         "支持家庭共享视角的聚合看板。")
    H3("5.7 AI 视觉识别")
    para("用户拍摄低碳场景（如公共交通票据、节能设备），端侧上传云端推理，返回结构化行为建议并回填至碳记录，"
         "降低手动录入成本。")

    # ---------------- 6 术语和缩略语 ----------------
    H1("6 术语和缩略语")
    terms = [
        ("OpenHarmony / HarmonyOS", "开放鸿蒙操作系统及其分布式能力底座"),
        ("ArkTS / ArkUI", "鸿蒙应用开发语言与声明式 UI 框架"),
        ("API 24", "本工程目标软件版本（compile/compatible SdkVersion=24）"),
        ("SoftBus（软总线）", "鸿蒙分布式软总线，提供近场发现与跨设备通信"),
        ("RuoYi-Vue", "若依前后端分离快速开发框架（本工程后端 3.8.7）"),
        ("REST / HTTPS", "端云通信的表征状态转移接口与加密传输协议"),
        ("Token（鉴权令牌）", "端侧调用服务端 API 的凭证，含 App-Token 鉴权"),
        ("Redis / MySQL", "内存缓存/会话存储 与 关系型业务数据库"),
        ("碳普惠", "将个人低碳行为量化并予激励的公共服务机制"),
        ("碳足迹 / 减排量", "行为产生的碳排放 与 相对基线减少的排放量"),
        ("碳惠券 / 核销码", "以积分兑换的权益券；PT- 前缀核销码用于线下核销"),
        ("液态玻璃（GlassLight）", "自研 135° 光照玻璃材质模型"),
        ("苹果动效（AppleMotion）", "自研统一时长/曲线/弹簧的动效语言"),
        ("深色换血", "深色模式下整体重建静态色值的主题机制"),
        ("注册闸门零帧锁死", "注册后主框架第一帧即锁交互、真实上屏才解锁"),
        ("命名空间隔离", "以账号为键隔离 Preferences 持久数据，支持多账户"),
    ]
    for k, v in terms:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1.4
        rk = p.add_run(k + "：")
        rk.font.bold = True
        rk.font.size = Pt(12)
        set_eastasia(rk, "宋体")
        rv = p.add_run(v)
        rv.font.size = Pt(12)
        set_eastasia(rv, "宋体")

    out = os.path.join(os.path.dirname(os.path.dirname(__file__)),
                       "一触青绿_设计方案_图文版.docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    out = build_docx()
    print("SAVED:", out)
