# -*- coding: utf-8 -*-
"""
生成《一触青绿 · UI 与液态玻璃专项讲解》Word 文档。
定位：零基础同学周五 8:30 汇报用，重点讲 UI 和液态玻璃（GlassLight），
含代码指读、时间节奏、提问预案、以及备用重点。
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

FONTS = {
    "body": "宋体",
    "heading": "黑体",
}
BODY_SIZE = Pt(12)          # 小四
HEADING_SIZE = Pt(15)
SUBHEAD_SIZE = Pt(13)


def set_run_font(run, name=FONTS["body"], size=BODY_SIZE, bold=False, color=None):
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), name)
    rfonts.set(qn('w:ascii'), name)
    rfonts.set(qn('w:hAnsi'), name)


def add_para(doc, text, size=BODY_SIZE, bold=False, name=FONTS["body"],
             align=None, color=None, first_indent=True, space_after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if first_indent and not align == WD_ALIGN_PARAGRAPH.CENTER:
        p.paragraph_format.first_line_indent = Pt(24)  # 首行缩进 2 字符
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_font(run, name=name, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    size = HEADING_SIZE if level == 1 else SUBHEAD_SIZE
    name = FONTS["heading"]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    set_run_font(run, name=name, size=size, bold=True)
    return p


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        r1 = p.add_run(bold_lead)
        set_run_font(r1, name=FONTS["body"], size=BODY_SIZE, bold=True)
        r2 = p.add_run(text)
        set_run_font(r2, name=FONTS["body"], size=BODY_SIZE)
    else:
        r = p.add_run(text)
        set_run_font(r, name=FONTS["body"], size=BODY_SIZE)
    return p


def code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Pt(18)
    for i, ln in enumerate(lines):
        run = p.add_run(ln + ("\n" if i < len(lines) - 1 else ""))
        set_run_font(run, name="Consolas", size=Pt(10.5), color=(40, 70, 120))
    return p


doc = Document()

# 全局正文默认字体
style = doc.styles["Normal"]
style.font.name = FONTS["body"]
style.font.size = BODY_SIZE
style.element.rPr.rFonts.set(qn('w:eastAsia'), FONTS["body"])

# ========================= 封面 =========================
add_para(doc, "一触青绿 · UI 与液态玻璃专项讲解",
         size=Pt(22), bold=True, name=FONTS["heading"],
         align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, space_after=4)
add_para(doc, "—— 周五 8:30 汇报专用（零基础可讲版）",
         size=Pt(13), name=FONTS["heading"],
         align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, space_after=18)
add_para(doc, "适用对象：对代码基本零基础、需要在老师面前讲清楚「本项目的 UI 设计与液态玻璃实现」的同学。",
         align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, color=(90, 90, 90))
add_para(doc, "说明：本文档所有代码位置均已对照工程实际文件核实（行号见各小节）。你只需要在讲的时候「指一指」对应文件即可，不必逐字背代码。",
         align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, color=(90, 90, 90))

# ========================= 一、汇报总览 =========================
add_heading(doc, "一、汇报总览（先给老师一句话定位）")
add_para(doc, "开场第一句就点题：「我们项目的 UI 最大亮点，是自己做了一套叫『液态玻璃（GlassLight）』的视觉材质，"
              "以及一套统一的苹果风动效（AppleMotion）。它不是套用现成组件库，而是从光照物理模型出发手写的，"
              "所以看起来通透、有立体感，而且深浅色都适配、低端机也不卡。」")
add_para(doc, "这句话有三个信息量：① 自研（不是抄的）；② 有物理依据（光照模型）；③ 工程上做了兼容（深浅色 + 画质降级）。"
              "后面所有内容都是围绕把这句话「撑满」。")
add_para(doc, "时间分配建议（10 分钟版）：液态玻璃 5 分钟 → 配套 UI 动效 2.5 分钟 → 代码指读 1.5 分钟 → 收尾 + 备用重点 1 分钟。"
              "如果只给 5 分钟，见最后一章「5 分钟精简路线」。")

# ========================= 二、液态玻璃 GlassLight 讲解主线 =========================
add_heading(doc, "二、液态玻璃 GlassLight 讲解主线（重点中的重点）")

add_heading(doc, "2.1 为什么不用现成的玻璃组件，要自己写？", level=2)
add_bullet(doc, "HarmonyOS 自带 .backdropBlur()（背景模糊）和通用组件，但只能做出「毛玻璃」的模糊，"
                "做不出「液态」的——也就是光从一处打过来、边缘有亮边、表面有高光流动那种通透立体感。")
add_bullet(doc, "所以项目用一个文件 GlassLight.ets（共 1303 行）把所有玻璃的「颜色、亮度、阴影、模糊、描边」算好，"
                "其它页面只管「我要一块 CARD 档玻璃」，不用关心光怎么打。这就是「一套材质，全应用复用」。")
add_bullet(doc, "证据：全工程 26 个文件都引用了 GlassLight，说明它不是孤芳自赏，而是真被当成基础设施在用。")

add_heading(doc, "2.2 核心思想：左上 45° 唯一主光源（用大白话讲）", level=2)
add_para(doc, "把它想成「真实世界里的一块玻璃板」：")
add_bullet(doc, "光从「左上角」照过来（45° 方向），所以玻璃的左上最亮、有高光；右下半部分稍微暗一点。")
add_bullet(doc, "玻璃边缘有一圈「菲涅尔亮边」——就像你斜看玻璃杯边缘会有一圈反光，这是玻璃「看起来像玻璃」的关键。")
add_bullet(doc, "表面还叠了一层很淡的「对角光染」（135°），让光有流动方向，而不是死板的均匀色。")
add_para(doc, "代码上，这套光照被压成一条「渐变（LinearGradient）」算出来，主入口是：")
code_block(doc, [
    "// ui/GlassLight.ets 第 902 行",
    "static surface(base, tier, ambient, gloss): LinearGradient { ... }",
    "// 用法（任意页面）：把一块 CARD 档玻璃的渐变直接套到背景上",
    "linearGradient(this.tw(GlassLight.surface(Glass.CARD_BG, GlassTier.CARD, '', 1)))",
])

add_heading(doc, "2.3 三个档位：pill < card < panel（面积越大越「收敛」）", level=2)
add_para(doc, "玻璃分三种尺寸档位（GlassTier，见 GlassLight.ets 第 126–130 行）：")
add_bullet(doc, "PILL（药丸/小标签，如分类 chip）—— 面积最小，可以最亮、最「湿」。")
add_bullet(doc, "CARD（卡片，如文章卡、图表卡）—— 中等。")
add_bullet(doc, "PANEL（大面板，如弹窗、底部栏）—— 面积最大，反而要最克制，否则一整块亮玻璃会刺眼。")
add_bullet(doc, "记住一句话：「面积越大，光越收着打」。这是一个反直觉但很专业的细节，老师听了会点头。")

add_heading(doc, "2.4 深色模式不是「反色」，而是「换血」", level=2)
add_para(doc, "很多 App 深色模式只是把浅色参数取反，结果玻璃在深色下很脏。本项目的做法是："
              "另写一整套深色参数表（washAlphaDark / edgeAlphaDark / innerAlphaDark，见 GlassLight.ets 第 513–563 行），"
              "深色下用专门的「暗部光源色 / 暗部阴影色」重新配（第 181–197 行 DARK_LIT / DARK_SHADE / DARK_TINT / DARK_SCRIM）。")
add_para(doc, "讲法：「我们的深色玻璃不是把浅色翻过来，而是根据暗环境下的光重新调了一遍，所以深色一样通透、不发黑。」")

add_heading(doc, "2.5 低端机不卡：三档画质降级", level=2)
add_bullet(doc, "GlassLight 内置画质档（GlassQualityLevel，第 142–144 行）：CINEMATIC（电影级，最漂亮）/ BALANCED（均衡）/ LITE（轻量）。")
add_bullet(doc, "默认开 CINEMATIC（第 200 行），但在低端机或检测到性能压力时，可以切到 LITE，减少模糊和层数，保证流畅。")
add_bullet(doc, "另外模糊有上限保护：blurOf()（第 739 行）/ blurScaled()（第 724 行）控制模糊半径，大卡片永不模糊（只用渐变模拟），避免性能崩。")
add_para(doc, "讲法：「漂亮和流畅我们都要——默认最漂亮，但留了降级开关，低端机自动变轻。」")

add_heading(doc, "2.6 滚动时玻璃会「跟着光走」：视差", level=2)
add_bullet(doc, "玻璃高光不是死贴在屏幕上，而是随页面滚动轻微平移，制造「光在玻璃里流动」的感觉。")
add_bullet(doc, "由 AppStorage 的 glassParallax（第 168 行）驱动，GlassLight.nudge(dy)（第 1245 行）更新，resetParallax()（第 1251 行）复位。")
add_para(doc, "讲法可以现场演示：上下滑一下首页，看卡片高光是不是微微在动——这就是视差。")

# ========================= 三、落地载体 GlassSurface =========================
add_heading(doc, "三、玻璃是怎么「贴」到界面上的：GlassSurface.ets")
add_para(doc, "GlassLight 只负责「算颜色」，真正把玻璃画出来的组件在 GlassSurface.ets（共 396 行）。它提供两个现成组件：")
add_bullet(doc, "GlassPanel（第 93 行）：内容面板，文章卡、Hero、弹窗、空态都用它。")
add_bullet(doc, "GlassPill（第 242 行）：小胶囊，分类标签 chip、收藏按钮、分段药丸都用它。")
add_para(doc, "它们内部都调 GlassLight 算好的结果：")
code_block(doc, [
    "// GlassSurface.ets 第 203–210 行（GlassPanel 内部）",
    ".border({ width: 1, color: this.tw(GlassLight.borderColors(this.tier, this.ambient, 1)) })",
    ".backdropBlur(GlassLight.blurOf(this.tier, this.glass))",
    ".shadow(this.tw(GlassLight.shadowOf(this.tier)))",
])
add_para(doc, "重点是第 38–39 行的 GlassLightLayers（玻璃光照层栈 Builder）：它把玻璃一层层叠出来——"
              "① 接触纱 → ② 对角光染（含内阴影）→ ③ 顶边菲涅尔亮边 → ④ 左边菲涅尔亮边 → ⑤ 内容。"
              "这正是 2.2 说的「左上主光源 + 菲涅尔边缘」在代码里的落点。")
add_para(doc, "一个工程坑（可以当彩蛋讲）：玻璃层原本用 Stack 堆叠，结果把父容器撑爆（卡片被拉成 210×586vp 的怪样子）。"
              "后来改成「零子层 + 绝对定位」方案（第 178–184 行注释记录的整改），才修好。说明「好看」背后是真踩过坑的。")

# ========================= 四、配套 UI 亮点：AppleMotion 动效 =========================
add_heading(doc, "四、配套 UI 亮点：统一的苹果风动效（AppleMotion.ets）")
add_para(doc, "液态玻璃解决「静态好看」，AppleMotion（共 229 行）解决「动起来也好看、且统一」。")
add_heading(doc, "4.1 统一时长档位（不让人觉得忽快忽慢）", level=2)
add_bullet(doc, "PRESS=100ms（按下反馈）、MICRO=200ms（小动效）、BASE=350ms（常规）、SHEET=500ms（弹窗）、THEME=380ms（换肤）。")
add_bullet(doc, "（见 AppleMotion.ets 第 84–92 行）所有动画都从这五档里挑，所以全应用节奏一致。")
add_heading(doc, "4.2 弹簧曲线：iOS 质感来源", level=2)
add_bullet(doc, "spring() / springSnappy() / springHeavy()（第 149–160 行）用 ArkUI 的 curves.springMotion 做「Q 弹回弹」。")
add_bullet(doc, "入场统一走 springLike()（第 134 行），预览器和真机两端行为一致，避免有的环境弹、有的环境不弹。")
add_heading(doc, "4.3 按压反馈 + 错峰入场", level=2)
add_bullet(doc, "按下时缩到 96%（PRESS_SCALE=0.96，第 115 行），给人「能戳下去」的真实感。")
add_bullet(doc, "首页看板 6 张卡片用 stagger(index)（第 226 行）「依次落座」，总时长 ≤300ms，既生动又不拖沓。")
add_bullet(doc, "转场有现成工厂：enterRise()（升起入场）、exitFade()（淡出）、stepIn/stepOut()（内容步进）、badgeIn()（角标弹入），第 170–212 行。")
add_para(doc, "讲法：「按钮、卡片、弹窗的动效我们统一了一套语言和节奏，所以整个 App 摸起来是一个整体，不会这块快那块慢。」")

# ========================= 五、代码指读（现场翻文件用） =========================
add_heading(doc, "五、代码指读清单（现场翻给老师看）")
add_para(doc, "提前把这几个文件在 DevEco Studio 里打开、翻到对应行，讲的时候直接指：")
add_bullet(doc, "GlassLight.ets 第 902 行 surface() —— 「这就是算一块玻璃渐变的核心函数」。")
add_bullet(doc, "GlassLight.ets 第 126–130 行 PILL/CARD/PANEL —— 「三种尺寸档位，面积越大光越收」。")
add_bullet(doc, "GlassLight.ets 第 142–144 行 CINEMATIC/BALANCED/LITE —— 「三档画质降级」。")
add_bullet(doc, "GlassLight.ets 第 513–563 行 *Dark 系列 —— 「深色模式是换血不是反色」。")
add_bullet(doc, "GlassSurface.ets 第 38–39 行 GlassLightLayers —— 「五层光栈，左上主光源 + 菲涅尔边」。")
add_bullet(doc, "GlassSurface.ets 第 93 / 242 行 GlassPanel / GlassPill —— 「业务页面只管要哪种玻璃」。")
add_bullet(doc, "ArticleListPage.ets 第 128 / 654 行 —— 「看，列表里文章卡用 GlassPanel，分类标签用 GlassPill，全应用 26 处复用」。")
add_bullet(doc, "AppleMotion.ets 第 84–92 行 —— 「五档时长，统一节奏」。")
add_bullet(doc, "AppleMotion.ets 第 149–160 / 226 行 —— 「弹簧曲线 + 错峰入场 stagger」。")

# ========================= 六、提问预案 =========================
add_heading(doc, "六、老师可能提问 & 你怎么答（预案）")
add_bullet(doc, "问「这玻璃和鸿蒙自带的毛玻璃有啥区别？」→ 答：自带只能模糊背景，我们额外做了「左上主光源 + 菲涅尔亮边 + 对角光染 + 滚动视差」，所以是液态通透感，不是死板模糊。")
add_bullet(doc, "问「深色模式怎么做的？」→ 答：不是反色，是另写一套深色光照参数表，暗环境下重新配光源色和阴影色，所以不发黑。")
add_bullet(doc, "问「低端机不会卡吗？」→ 答：有 CINEMATIC/BALANCED/LITE 三档，大卡片永不模糊只用渐变模拟，并限制模糊半径。")
add_bullet(doc, "问「这是你们自己写的还是抄的？」→ 答：自己写的，一个 1303 行的 GlassLight 文件，被 26 个页面复用，是项目自研的视觉基础设施。")
add_bullet(doc, "问「动效为什么统一？」→ 答：AppleMotion 定了五档时长和弹簧曲线，所有按钮/卡片/弹窗都从里面取，所以全 App 手感一致。")

# ========================= 七、备用重点（万一时间富余或被问业务） =========================
add_heading(doc, "七、备用重点（时间富余 / 被问业务时拿出来）")
add_para(doc, "如果老师想多听点「系统本身在干嘛」，准备这三条，每条一句话就能讲清：")
add_bullet(doc, "① 减排算账（业务逻辑主线）：记录一次低碳行为 → 减排量(kg)=因子×数量 → 积分=kg×10 → 200 积分换 1 棵树/1 张证书。"
                "代码在 data/LocalStore.ets 的 addRecord() 和 ui/Theme.ets 的 Biz.pointsOf()。")
add_bullet(doc, "② 碳惠券：用积分兑换券，券有 7 天有效期（从领取起算），核销时生成一张 ASCII 单行二维码（约 146 字节），可幂等回看。")
add_bullet(doc, "③ 多账号隔离：每个账号的数据用「业务键@账号」分开存，切换账号互不串数据；注册闸门用零帧锁死，防止页面没真正上屏就被跳过。")
add_para(doc, "（以上三条在《一触青绿_业务逻辑讲解稿.docx》里有零基础详细版，可搭配使用。）")

# ========================= 八、5 分钟精简路线 =========================
add_heading(doc, "八、如果只给 5 分钟：精简路线")
add_bullet(doc, "0:00–0:30 开场定位（一句话：自研液态玻璃 + 统一动效，有物理模型、深浅色适配、低端不卡）。")
add_bullet(doc, "0:30–2:30 液态玻璃核心：左上 45° 主光源 + 菲涅尔亮边 + 三档尺寸（pill<card<panel）+ 深色换血 + 三档画质。指 GlassLight.ets 第 902 / 126 / 142 / 513 行。")
add_bullet(doc, "2:30–3:30 落地：GlassPanel / GlassPill 两个组件，26 处复用；指 GlassSurface.ets 第 38 / 93 / 242 行。")
add_bullet(doc, "3:30–4:30 动效：AppleMotion 五档时长 + 弹簧 + 按压 96% + stagger；指 AppleMotion.ets 第 84 / 149 / 226 行。")
add_bullet(doc, "4:30–5:00 收尾：强调自研、体系化、工程上做了兼容；若有空补一句减排算账（备用重点①）。")

add_para(doc, "—— 文档结束。祝你周五汇报顺利。", align=WD_ALIGN_PARAGRAPH.CENTER,
         first_indent=False, color=(90, 90, 90), space_after=2)

out = "D:/11/carbon-master/一触青绿_UI与液态玻璃_专项讲解.docx"
doc.save(out)
print("saved:", out)
