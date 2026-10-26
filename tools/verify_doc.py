# -*- coding: utf-8 -*-
"""校验生成的 docx：字体、行距、表格、图数量、禁用词扫描"""
import glob
from docx import Document
from docx.shared import Pt

path = glob.glob('tools/一触青绿*.docx')[0]
doc = Document(path)

print('文件:', path)
import os
print('大小(KB): %.1f' % (os.path.getsize(path) / 1024))
print('段落数:', len(doc.paragraphs))
print('表格数:', len(doc.tables))

# Normal 样式检查
st = doc.styles['Normal']
print('Normal 字体:', st.font.name, '| 字号:', st.font.size.pt if st.font.size else None)
print('Normal 行距:', st.paragraph_format.line_spacing)

# 抽查正文 run 字体
sample_p = None
for p in doc.paragraphs:
    if p.text.startswith('一触青绿是一款'):
        sample_p = p
        break
if sample_p is not None and sample_p.runs:
    sample = sample_p.runs[0]
    from docx.oxml.ns import qn
    rPr = sample._element.rPr
    rF = rPr.find(qn('w:rFonts'))
    print('正文 run eastAsia:', rF.get(qn('w:eastAsia')), '| 字号:', sample.font.size.pt,
          '| 段落行距:', sample_p.paragraph_format.line_spacing)

# 全文文本（含表格）
texts = [p.text for p in doc.paragraphs]
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            texts.append(c.text)
full = '\n'.join(texts)
print('总字符数:', len(full))

# Mermaid 图计数
import re
figs = re.findall(r'flowchart [LT]B|flowchart LR', full)
print('Mermaid 图数量:', len(figs))

# 禁用词扫描（应仅出现在否定性边界声明中）
for kw in ['http', '服务器', '云端', '联网', '多端', '分布式', '组网', '软总线', '账号同步']:
    hits = [l.strip()[:60] for l in full.split('\n') if kw in l]
    print(f'[{kw}] 命中 {len(hits)} 处')
    for h in hits:
        print('   ·', h)
